#!/usr/bin/env python3
"""Build a simple bill-split .xlsx from a JSON config.

Reads JSON from stdin (or a file path arg for testing with fake data).
Never persist a config containing real receipt data; pipe it in.

    .venv/bin/python build_bill_split.py <<'JSON'
    { ... }
    JSON
    .venv/bin/python build_bill_split.py fixtures/fake.json  # fake data only

The config shape is defined and validated against schema.json sitting next
to this script. Items without "assigned" are left blank (red highlight).
"""

import datetime
import json
import re
import sys
from pathlib import Path
from typing import NotRequired, TypedDict

from jsonschema import FormatChecker, validate
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

MONEY = '"$"#,##0.00'
SCHEMA_PATH = Path(__file__).resolve().parent / "schema.json"
_BAD_TITLE = re.compile(r"[:\\/?*\[\]]")  # Excel forbids these in sheet names


class Item(TypedDict):
    name: str
    price: float
    assigned: NotRequired[str]


class Receipt(TypedDict):
    title: str
    tax: float
    tip: float
    subtotal: float
    items: list[Item]
    discount: NotRequired[float]


class Config(TypedDict):
    people: list[str]
    filename: str
    receipts: list[Receipt]
    date: NotRequired[str]


def sheet_title(title: str, used: set[str]) -> str:
    """Excel tab name: ≤31 chars, no illegal chars, unique ignoring case."""
    cleaned = " ".join(_BAD_TITLE.sub("", title).split()) or "receipt"
    cleaned = cleaned[:31]
    if cleaned.lower() in {"history", "_people"}:
        cleaned = "receipt"
    name = cleaned
    n = 2
    taken = {u.lower() for u in used}
    while name.lower() in taken:
        suffix = f" ({n})"
        name = cleaned[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def excel_quote(s: str) -> str:
    """Escape " for use inside an Excel formula string."""
    return s.replace('"', '""')


def make_sheet(
    ws: Worksheet,
    items: list[Item],
    people: list[str],
    tax: float,
    tip: float,
    receipt_subtotal: float,
    title: str,
    date: str,
    discount: float = 0,
) -> None:
    """Item | Price | Assigned To, with a per-person summary below."""
    n = len(people)
    # ponytail: park receipt totals past the last person column (and past C)
    # so groups of 5+ don't overwrite F1:G5.
    label_col = max(6, n + 3)
    value_col = label_col + 1
    value_letter = get_column_letter(value_col)

    ws["A1"] = f"{title} — {date}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

    # Receipt totals (Total is a formula so editing tax/tip updates it)
    labels = ["Subtotal", "Tax", "Tip", "Discount", "Total"]
    values: list[float | str] = [
        receipt_subtotal,
        tax,
        tip,
        discount,
        f"={value_letter}1+{value_letter}2+{value_letter}3-{value_letter}4",
    ]
    for row, (label, value) in enumerate(zip(labels, values), 1):
        ws.cell(row=row, column=label_col, value=label)
        cell = ws.cell(row=row, column=value_col, value=value)
        cell.number_format = MONEY

    for col, header in enumerate(["Item", "Price", "Assigned To"], 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    first_data_row = 3
    for i, item in enumerate(items):
        row = first_data_row + i
        ws.cell(row=row, column=1, value=item["name"])
        ws.cell(row=row, column=2, value=item["price"]).number_format = MONEY
        ws.cell(row=row, column=3, value=item.get("assigned", ""))

    last_data_row = first_data_row + len(items) - 1

    # Dropdown from hidden _people sheet (range, so commas in names are fine)
    dv = DataValidation(
        type="list",
        formula1=f"=_people!$A$1:$A${n}",
        allow_blank=True,
    )
    dv.error = "Pick a person."
    dv.errorTitle = "Invalid"
    ws.add_data_validation(dv)
    dv.add(f"C{first_data_row}:C{last_data_row}")

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006")
    ws.conditional_formatting.add(
        f"A{first_data_row}:C{last_data_row}",
        FormulaRule(
            formula=[f'$C{first_data_row}=""'],
            fill=red_fill,
            font=red_font,
        ),
    )

    # Per-person summary: SUMIF subtotals, then tax/tip/discount by share
    summary_row = last_data_row + 2
    ws.cell(row=summary_row, column=1, value="Person").font = Font(bold=True)
    for p_idx, person in enumerate(people, 2):
        ws.cell(row=summary_row, column=p_idx, value=person).font = Font(bold=True)

    for offset, label in enumerate(["Subtotal", "Tax", "Tip", "Discount", "Total Owed"], 1):
        ws.cell(row=summary_row + offset, column=1, value=label).font = Font(bold=True)

    subtotal_row = summary_row + 1
    tax_row = summary_row + 2
    tip_row = summary_row + 3
    discount_row = summary_row + 4
    total_row = summary_row + 5
    last_person_col = get_column_letter(1 + n)

    for p_idx, person in enumerate(people, 2):
        col_letter = get_column_letter(p_idx)
        quoted = excel_quote(person)
        share = (
            f"IFERROR({col_letter}{subtotal_row}"
            f"/SUM($B${subtotal_row}:${last_person_col}{subtotal_row})"
        )
        ws.cell(
            row=subtotal_row,
            column=p_idx,
            value=(
                f"=SUMIF($C${first_data_row}:$C${last_data_row},"
                f'"{quoted}",$B${first_data_row}:$B${last_data_row})'
            ),
        ).number_format = MONEY
        ws.cell(
            row=tax_row,
            column=p_idx,
            value=f"={share}*${value_letter}$2,0)",
        ).number_format = MONEY
        ws.cell(
            row=tip_row,
            column=p_idx,
            value=f"={share}*${value_letter}$3,0)",
        ).number_format = MONEY
        ws.cell(
            row=discount_row,
            column=p_idx,
            value=f"={share}*${value_letter}$4,0)",
        ).number_format = MONEY
        total_cell = ws.cell(
            row=total_row,
            column=p_idx,
            value=(
                f"={col_letter}{subtotal_row}+{col_letter}{tax_row}"
                f"+{col_letter}{tip_row}-{col_letter}{discount_row}"
            ),
        )
        total_cell.number_format = MONEY
        total_cell.font = Font(bold=True)

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 12
    for col in range(2, 2 + n):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.column_dimensions[get_column_letter(label_col)].width = 12
    ws.column_dimensions[get_column_letter(value_col)].width = 12

    ws.freeze_panes = "A3"


def load_schema() -> dict:
    with open(SCHEMA_PATH) as f:
        return json.load(f)


def validate_config(config: Config) -> None:
    validate(instance=config, schema=load_schema(), format_checker=FormatChecker())
    # ponytail: "assigned must be in people" is a cross-field rule JSON Schema
    # can't express without non-standard $data refs; checked here instead.
    people = set(config["people"])
    for receipt in config["receipts"]:
        for item in receipt["items"]:
            assigned = item.get("assigned")
            if assigned is not None and assigned not in people:
                raise ValueError(
                    f"{receipt['title']}: item {item['name']!r} assigned to "
                    f"{assigned!r}, who is not in people {sorted(people)}"
                )


def build(config: Config) -> None:
    validate_config(config)
    people = config["people"]
    filename = config["filename"]
    date = config.get("date", datetime.date.today().isoformat())

    wb = Workbook()
    people_ws = wb.create_sheet("_people")
    for i, person in enumerate(people, 1):
        people_ws.cell(row=i, column=1, value=person)
    people_ws.sheet_state = "hidden"

    used_titles = {"_people"}
    first = True
    for receipt in config["receipts"]:
        items = receipt["items"]
        expected = receipt["subtotal"]
        calculated = sum(item["price"] for item in items)
        if round(calculated, 2) != round(expected, 2):
            raise ValueError(
                f"{receipt['title']} subtotal mismatch: "
                f"calculated ${calculated:.2f}, receipt ${expected:.2f}"
            )

        name = sheet_title(receipt["title"], used_titles)
        if first:
            ws = wb.active
            assert ws is not None
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(title=name)

        make_sheet(
            ws,
            items,
            people,
            tax=receipt["tax"],
            tip=receipt["tip"],
            receipt_subtotal=receipt["subtotal"],
            title=receipt["title"],
            date=date,
            discount=receipt.get("discount", 0),
        )

    wb.save(filename)
    print(f"Saved {filename}")


if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] == "-":
        raw = sys.stdin.read()
    else:
        with open(sys.argv[1]) as f:
            raw = f.read()
    build(json.loads(raw))
